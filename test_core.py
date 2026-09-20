import unittest
import tempfile
import threading
import time
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch

import requests
from openpyxl import Workbook, load_workbook

from app import app, available_port, clients, clients_lock
from coresys import (
    CoresysClient, extract_sla_info, format_awb_text, normalize_awbs,
    normalize_awb_targets, parse_tracking_focus_html, parse_tracking_history,
    split_awbs, split_date_range,
)
from history_export import (
    export_tracking_history, first_pod_courier, pickup_attempt_summary,
    read_awb_targets_workbook, tracking_milestones,
)
from jobs import JobManager
from tracking_focus_export import export_tracking_focus, read_reference_targets_workbook


class DateBatchTests(unittest.TestCase):
    def test_inclusive_batches(self):
        batches = split_date_range("2026-01-01", "2026-01-17", 7)
        self.assertEqual(
            [batch.as_dict() for batch in batches],
            [
                {"from": "2026-01-01", "to": "2026-01-07"},
                {"from": "2026-01-08", "to": "2026-01-14"},
                {"from": "2026-01-15", "to": "2026-01-17"},
            ],
        )

    def test_month_boundary_clamping(self):
        batches = split_date_range("2026-01-28", "2026-02-05", 7)
        self.assertEqual(
            [batch.as_dict() for batch in batches],
            [
                {"from": "2026-01-28", "to": "2026-01-31"},
                {"from": "2026-02-01", "to": "2026-02-05"},
            ],
        )

    def test_30_day_month_clamping(self):
        batches = split_date_range("2026-04-29", "2026-05-04", 7)
        self.assertEqual(
            [batch.as_dict() for batch in batches],
            [
                {"from": "2026-04-29", "to": "2026-04-30"},
                {"from": "2026-05-01", "to": "2026-05-04"},
            ],
        )

    def test_rejects_reverse_range(self):
        with self.assertRaises(ValueError):
            split_date_range("2026-02-02", "2026-02-01", 7)


class PortTests(unittest.TestCase):
    def test_skips_an_occupied_port(self):
        import socket

        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as listener:
            listener.bind(("127.0.0.1", 0))
            occupied = int(listener.getsockname()[1])
            self.assertNotEqual(available_port(occupied, attempts=1), occupied)


class AwbBatchTests(unittest.TestCase):
    def test_normalizes_and_deduplicates(self):
        self.assertEqual(normalize_awbs("cgk1\nCGK2, cgk1;CGK3"), ["CGK1", "CGK2", "CGK3"])

    def test_splits_at_limit(self):
        batches = split_awbs([f"AWB{i}" for i in range(20_001)])
        self.assertEqual([len(batch) for batch in batches], [10_000, 10_000, 1])

    def test_report_payload_uses_browser_crlf_line_endings(self):
        self.assertEqual(format_awb_text(["AWB1", "AWB2"]), "AWB1\r\nAWB2")

    def test_parses_awb_and_tlc_columns(self):
        raw = "No. AWB\tTLC Tujuan\nCGK1600227657333\tPKY\nPKN1600227646503 PKY"
        self.assertEqual(normalize_awb_targets(raw), [
            {"awb": "CGK1600227657333", "tlc": "PKY"},
            {"awb": "PKN1600227646503", "tlc": "PKY"},
        ])


class TrackingHistoryTests(unittest.TestCase):
    HISTORY_HTML = """
        <h4 class="form-section">Riwayat</h4>
        <table><thead><tr><th>#</th><th>Proses</th><th>No. Dokumen</th>
        <th>Referensi / Tag</th><th>Tanggal &amp; Jam</th><th>Lokasi / Oleh</th>
        <th>Keterangan</th></tr></thead><tbody><tr><td>1</td><td>POD</td>
        <td>DOC1</td><td>REF1</td><td>2026-07-22 10:30:00</td>
        <td>KANTOR PUSAT / USER</td><td>DELIVERED</td></tr></tbody></table>
    """

    def test_parses_history_table(self):
        self.assertEqual(parse_tracking_history(self.HISTORY_HTML), [{
            "no": 1,
            "process": "POD",
            "document": "DOC1",
            "reference": "REF1",
            "datetime": "2026-07-22 10:30:00",
            "location": "KANTOR PUSAT / USER",
            "note": "DELIVERED",
        }])

    def test_reuses_recent_tracking_history(self):
        client = CoresysClient()
        client.username = "TEST"
        response = Mock()
        response.url = "https://online.coresyssap.com/tracking/getriwayat"
        response.text = self.HISTORY_HTML
        response.raise_for_status.return_value = None
        client.session.post = Mock(return_value=response)
        self.assertEqual(len(client.fetch_tracking_history("AWB1")), 1)
        self.assertEqual(len(client.fetch_tracking_history("AWB1")), 1)
        self.assertEqual(client.session.post.call_count, 1)

    def test_selects_first_milestones_at_destination_tlc(self):
        records = [
            {"process": "ENTRI VERIFIED", "datetime": "2026-06-23 13:31:27", "location": "KANTOR PUSAT / CGK1"},
            {"process": "OUTGOING SMU", "datetime": "2026-06-25 06:48:35", "location": "KANTOR PUSAT / CGK1"},
            {"process": "INCOMING SMU", "datetime": "2026-06-26 03:04:33", "location": "SURABAYA / SUB1"},
            {"process": "INCOMING SMU", "datetime": "2026-07-02 13:56:28", "location": "PALANGKARAYA / PKY06208165A"},
            {"process": "POD", "datetime": "2026-07-03 09:31:00", "location": "PALANGKARAYA / PKYD103184915"},
        ]
        verified, outgoing, incoming, pod = tracking_milestones(records, "PKY")
        self.assertEqual(verified.strftime("%Y-%m-%d %H:%M:%S"), "2026-06-23 13:31:27")
        self.assertEqual(outgoing.strftime("%Y-%m-%d %H:%M:%S"), "2026-06-25 06:48:35")
        self.assertEqual(incoming.strftime("%Y-%m-%d %H:%M:%S"), "2026-07-02 13:56:28")
        self.assertEqual(pod.strftime("%Y-%m-%d %H:%M:%S"), "2026-07-03 09:31:00")

    def test_selects_courier_after_last_slash_from_first_pod(self):
        records = [
            {"process": "INCOMING SMU", "location": "PALANGKARAYA / PKY06208165A"},
            {"process": "POD", "location": "KANTOR CABANG PAMEKASAN / PMKD2112416758"},
            {"process": "POD", "location": "PALANGKARAYA / PKYD103184915"},
        ]
        self.assertEqual(first_pod_courier(records), "PMKD2112416758")

    def test_verifies_not_ready_pickup_attempt_within_h_plus_one(self):
        records = [
            {"process": "ENTRI (SEDANG DI PICKUP)", "datetime": "2026-08-08 15:34:40", "note": ""},
            {"process": "ENTRI (PENDING PICKUP)", "datetime": "2026-08-09 20:37:00",
             "note": "[KURIR: SOFIANDI] [KETERANGAN: PAKET BELUM READY - CUSTOMER]"},
            {"process": "PICKED UP", "datetime": "2026-08-13 17:46:08", "note": ""},
        ]
        result = pickup_attempt_summary(records)
        self.assertEqual(result[2:6], ["SOFIANDI", "PAKET BELUM READY - CUSTOMER", 1, 1])
        self.assertEqual(result[7], "SESUAI (H+1)")
        self.assertEqual(result[8], "SUDAH ADA PERCOBAAN PICKUP - PAKET BELUM READY")

    def test_marks_first_attempt_after_h_plus_one_as_late_pickup(self):
        records = [
            {"process": "ENTRI (SEDANG DI PICKUP)", "datetime": "2026-08-08 15:34:40", "note": ""},
            {"process": "ENTRI (PENDING PICKUP)", "datetime": "2026-08-10 20:37:00",
             "note": "[KURIR: SOFIANDI] [KETERANGAN: PAKET BELUM READY]"},
        ]
        result = pickup_attempt_summary(records)
        self.assertEqual(result[7], "LATE PICKUP (H+2)")
        self.assertEqual(result[8], "LATE PICKUP - SUDAH ADA PERCOBAAN PICKUP - PAKET BELUM READY")

    def test_marks_missing_attempt_past_h_plus_one_as_late_pickup(self):
        records = [{"process": "ENTRI (SEDANG DI PICKUP)", "datetime": "2020-01-01 10:00:00", "note": ""}]
        result = pickup_attempt_summary(records)
        self.assertEqual(result[7], "LATE PICKUP - BELUM ADA PERCOBAAN")
        self.assertEqual(result[8], "LATE PICKUP - BELUM ADA HASIL PERCOBAAN")

    def test_reads_awb_targets_from_excel_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "targets.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Catatan", None])
            sheet.append(["No. AWB", "TLC Tujuan"])
            sheet.append(["CGK1600227657333", "PKY"])
            sheet.append(["PKN1600227646503", "PKY"])
            workbook.save(path)
            workbook.close()
            with path.open("rb") as source:
                self.assertEqual(read_awb_targets_workbook(source), [
                    {"awb": "CGK1600227657333", "tlc": "PKY"},
                    {"awb": "PKN1600227646503", "tlc": "PKY"},
                ])

    def test_reads_awb_only_excel_for_courier_mode(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "awbs.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Daftar kiriman"])
            sheet.append(["No. AWB"])
            sheet.append(["CGK1600227657333"])
            sheet.append(["PKN1600227646503"])
            workbook.save(path)
            workbook.close()
            with path.open("rb") as source:
                self.assertEqual(read_awb_targets_workbook(source, require_tlc=False), [
                    {"awb": "CGK1600227657333", "tlc": ""},
                    {"awb": "PKN1600227646503", "tlc": ""},
                ])

    class FakeClient:
        def fork(self):
            return self

        def fetch_tracking_history(self, awb):
            if awb == "MISSING":
                return []
            return [{
                "no": 1,
                "process": "POD",
                "document": "DOC1",
                "reference": awb,
                "datetime": "2026-07-22 10:30:00",
                "location": "KANTOR PUSAT / USER",
                "note": "DELIVERED",
            }]

    def test_exports_all_history_to_one_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "history.xlsx"
            progress = []
            result, found, missing, failed = export_tracking_history(
                self.FakeClient(), ["AWB1", "MISSING"], output,
                progress=lambda done, total: progress.append((done, total)),
            )
            workbook = load_workbook(result, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["Milestone Tujuan", "Ringkasan", "Semua History"])
                self.assertEqual((found, missing, failed), (1, 1, 0))
                self.assertEqual(sum(1 for _ in workbook["Semua History"].iter_rows(values_only=True)), 2)
                self.assertEqual(sum(1 for _ in workbook["Milestone Tujuan"].iter_rows(values_only=True)), 3)
                statuses = {row[0]: row[6] for row in workbook["Ringkasan"].iter_rows(min_row=5, values_only=True)}
                self.assertEqual(statuses, {"AWB1": "BERHASIL", "MISSING": "TIDAK DITEMUKAN"})
                self.assertEqual(progress[-1], (2, 2))
            finally:
                workbook.close()

    def test_exports_milestone_only(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "milestone.xlsx"
            result, found, missing, failed = export_tracking_history(
                self.FakeClient(), [{"awb": "AWB1", "tlc": "PKY"}], output,
                include_summary=False, include_history=False,
            )
            workbook = load_workbook(result, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["Milestone Tujuan"])
                self.assertEqual((found, missing, failed), (1, 0, 0))
            finally:
                workbook.close()

    def test_exports_first_pod_courier_only(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "courier.xlsx"
            result, found, missing, failed = export_tracking_history(
                self.FakeClient(), [{"awb": "AWB1", "tlc": ""}], output,
                include_summary=False, include_history=False, tracking_mode="courier_pod",
            )
            workbook = load_workbook(result, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["POD Pertama"])
                rows = list(workbook["POD Pertama"].iter_rows(values_only=True))
                self.assertEqual(rows, [("No. AWB", "ID Kurir POD Pertama"), ("AWB1", "USER")])
                self.assertEqual((found, missing, failed), (1, 0, 0))
            finally:
                workbook.close()

    def test_exports_pickup_attempt_verification(self):
        class PickupClient(self.FakeClient):
            def fetch_tracking_history(self, awb):
                return [
                    {"no": 1, "process": "ENTRI (SEDANG DI PICKUP)", "document": "PUP1",
                     "reference": awb, "datetime": "2026-08-08 15:34:40", "location": "CABANG / USER", "note": ""},
                    {"no": 2, "process": "ENTRI (PENDING PICKUP)", "document": "PRF1",
                     "reference": "PUP1", "datetime": "2026-08-08 20:37:00", "location": "CABANG / ID - SOFIANDI",
                     "note": "[KURIR: SOFIANDI] [KETERANGAN: PAKET BELUM READY - CUSTOMER]"},
                ]

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "pickup.xlsx"
            result, *_ = export_tracking_history(
                PickupClient(), ["AWB1"], output, tracking_mode="pickup_attempt",
                include_summary=False, include_history=False,
            )
            workbook = load_workbook(result, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["Verifikasi Pickup"])
                row = next(workbook["Verifikasi Pickup"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(row[0], "AWB1")
                self.assertEqual(row[3:7], ("SOFIANDI", "PAKET BELUM READY - CUSTOMER", 1, 1))
                self.assertEqual(row[8:], ("SESUAI (H+0)", "SUDAH ADA PERCOBAAN PICKUP - PAKET BELUM READY"))
            finally:
                workbook.close()

    def test_tracking_export_allows_twelve_parallel_requests(self):
        class ConcurrentClient:
            def __init__(self, state=None):
                self.state = state or {"active": 0, "maximum": 0, "lock": threading.Lock()}

            def fork(self):
                return self.__class__(self.state)

            def fetch_tracking_history(self, awb):
                with self.state["lock"]:
                    self.state["active"] += 1
                    self.state["maximum"] = max(self.state["maximum"], self.state["active"])
                try:
                    time.sleep(0.1)
                    return [{
                        "no": 1, "process": "POD", "document": "", "reference": awb,
                        "datetime": "2026-07-22 10:30:00", "location": "CABANG / USER", "note": "",
                    }]
                finally:
                    with self.state["lock"]:
                        self.state["active"] -= 1

        with tempfile.TemporaryDirectory() as directory:
            client = ConcurrentClient()
            export_tracking_history(
                client, [f"AWB{i}" for i in range(12)], Path(directory) / "parallel.xlsx",
                parallelism=12, include_summary=False, include_history=False,
                tracking_mode="courier_pod",
            )
            self.assertEqual(client.state["maximum"], 12)


class UrlTests(unittest.TestCase):
    def test_pickup_url_matches_portal_contract(self):
        client = CoresysClient()
        url = client.pickup_url(
            "2026-01-01",
            "2026-01-07",
            {"date_pickup": "1", "pilih_status": "0"},
            "report_monitoring_xlsx",
        )
        self.assertIn("/pickup/report_monitoring_xlsx/-/01-01-2026/07-01-2026/0/", url)
        self.assertTrue(url.endswith("?token=token_02"))

    def test_pickup_manual_url_matches_current_portal_contract(self):
        client = CoresysClient()
        url = client.pickup_manual_url(
            "2026-01-01",
            "2026-01-01",
            {
                "pilih_status": "-",
                "date_pickup": "0",
                "origin_area_branch_code": "1",
                "destination_area_branch_code": "-",
            },
            "report_monitoring",
        )
        self.assertIn("/pickup_manual/report_monitoring/-/01-01-2026/01-01-2026/-/", url)
        self.assertTrue(url.endswith("/-/1/-"))


class XlsxResultTests(unittest.TestCase):
    def test_detects_header_only_awb_result(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.xlsx"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("xl/worksheets/sheet1.xml", b"<sheet><row/><row/><row/></sheet>")
            self.assertEqual(JobManager._xlsx_data_rows(path, header_rows=3), 0)


class ParallelJobTests(unittest.TestCase):
    class FakeClient:
        def __init__(self, state=None):
            self.state = state if state is not None else {"active": 0, "maximum": 0, "lock": threading.Lock(), "timestamps": []}

        def fork(self):
            return self.__class__(self.state)

        def pickup_url(self, start, end, filters, export):
            return f"https://example.test/{start}/{end}"

        def download_direct(self, url, destination, progress):
            with self.state["lock"]:
                self.state["active"] += 1
                self.state["maximum"] = max(self.state["maximum"], self.state["active"])
                self.state.setdefault("timestamps", []).append(time.monotonic())
            try:
                time.sleep(0.1)
                destination.parent.mkdir(parents=True, exist_ok=True)
                destination.write_bytes(b"result")
                progress(6, 6)
                return destination
            finally:
                with self.state["lock"]:
                    self.state["active"] -= 1

    def test_runs_batches_in_parallel(self):
        with tempfile.TemporaryDirectory() as directory:
            manager = JobManager(Path(directory), max_concurrent_batches=3)
            client = self.FakeClient()
            try:
                job = manager.create(client, {
                    "workflow": "pickup",
                    "export": "report_monitoring_xlsx",
                    "date_from": "2026-01-01",
                    "date_to": "2026-01-03",
                    "batch_days": 1,
                    "delay_seconds": 0,
                    "parallelism": 2,
                })
                deadline = time.monotonic() + 5
                while time.monotonic() < deadline:
                    current = manager.public(job["id"])
                    if current["status"] not in {"queued", "running"}:
                        break
                    time.sleep(0.02)
                self.assertEqual(current["status"], "complete")
                self.assertEqual(current["completed"], 3)
                self.assertEqual(client.state["maximum"], 2)
            finally:
                manager.shutdown()

    def test_runs_batches_sequentially_with_delay(self):
        with tempfile.TemporaryDirectory() as directory:
            manager = JobManager(Path(directory), max_concurrent_batches=1)
            client = self.FakeClient()
            timestamps = []

            try:
                job = manager.create(client, {
                    "workflow": "pickup",
                    "export": "report_monitoring_xlsx",
                    "date_from": "2026-01-01",
                    "date_to": "2026-01-02",
                    "batch_days": 1,
                    "delay_seconds": 1,
                    "parallelism": 1,
                })
                deadline = time.monotonic() + 5
                while time.monotonic() < deadline:
                    current = manager.public(job["id"])
                    if current["status"] not in {"queued", "running"}:
                        break
                    time.sleep(0.02)
                self.assertEqual(current["status"], "complete")
                self.assertEqual(current["completed"], 2)
                self.assertEqual(client.state["maximum"], 1)
                timestamps = client.state["timestamps"]
                self.assertEqual(len(timestamps), 2)
                self.assertGreaterEqual(timestamps[1] - timestamps[0], 1.0)
            finally:
                manager.shutdown()


class JobRetryTests(unittest.TestCase):
    class FlakyClient:
        def __init__(self):
            self.attempts = 0
            self.lock = threading.Lock()

        def fork(self):
            return self

        def pickup_url(self, start, end, filters, export):
            return f"https://example.test/{start}/{end}"

        def download_direct(self, url, destination, progress):
            with self.lock:
                self.attempts += 1
                should_fail = (self.attempts == 1)
            if should_fail:
                raise RuntimeError("Koneksi portal terputus.")
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_bytes(b"ok")
            progress(2, 2)
            return destination

    def test_retry_failed_batch(self):
        with tempfile.TemporaryDirectory() as directory:
            manager = JobManager(Path(directory), max_concurrent_batches=1)
            client = self.FlakyClient()
            try:
                job = manager.create(client, {
                    "workflow": "pickup",
                    "export": "report_monitoring_xlsx",
                    "date_from": "2026-01-01",
                    "date_to": "2026-01-01",
                    "batch_days": 1,
                    "delay_seconds": 0,
                    "parallelism": 1,
                })
                deadline = time.monotonic() + 5
                while time.monotonic() < deadline:
                    cur = manager.public(job["id"])
                    if cur["status"] in {"complete_with_errors", "failed"}:
                        break
                    time.sleep(0.02)
                self.assertEqual(cur["status"], "complete_with_errors")
                self.assertEqual(cur["batches"][0]["status"], "failed")
                self.assertEqual(cur["failed"], 1)

                # Retry the failed batch
                manager.retry_batch(job["id"], 1, client)
                deadline = time.monotonic() + 5
                while time.monotonic() < deadline:
                    cur = manager.public(job["id"])
                    if cur["status"] == "complete":
                        break
                    time.sleep(0.02)
                self.assertEqual(cur["status"], "complete")
                self.assertEqual(cur["batches"][0]["status"], "complete")
                self.assertEqual(cur["completed"], 1)
                self.assertEqual(cur["failed"], 0)
            finally:
                manager.shutdown()

    def test_retry_batch_api_endpoint(self):
        with app.test_client() as tc:
            c = CoresysClient()
            c.username = "TEST_USER"
            with tc.session_transaction() as sess:
                sess["session_id"] = "test-sess-retry"
            with clients_lock:
                clients["test-sess-retry"] = c

            res = tc.post("/api/jobs/non-existent/batches/1/retry")
            self.assertEqual(res.status_code, 404)


class PodPollingTests(unittest.TestCase):
    @patch("coresys.time.sleep", return_value=None)
    def test_retries_transient_disconnect(self, _sleep):
        client = CoresysClient()
        processing = Mock()
        processing.raise_for_status.return_value = None
        processing.json.return_value = {"status": 1}
        client.session.post = Mock(side_effect=[requests.ConnectionError("closed"), processing])

        self.assertEqual(client.poll_pod_v2("123"), {"state": "processing"})
        self.assertEqual(client.session.post.call_count, 2)


class TrackingFocusTests(unittest.TestCase):
    SAMPLE_HTML = """
        <table class="table view_sample">
            <thead>
                <tr>
                    <th>No</th><th>No. AWB</th><th>No. Referance</th><th>No. Master</th>
                    <th>Tanggal</th><th>Asal</th><th>District Asal</th><th>Tujuan</th>
                    <th>Kota Tujuan</th><th>Jenis Layanan</th><th>Transaksi</th><th>Kilo</th>
                    <th>Koli</th><th>No Invoice</th><th>Status</th><th>Detail Status</th>
                    <th>No. Resi Pickup</th><th>Dibuat Oleh</th><th>Di POD Oleh</th><th>Riwayat</th>
                </tr>
            </thead>
            <tbody>
                <tr>
                    <td>1</td>
                    <td><a href="...">CGK8161744700066</a></td>
                    <td>NONPO03072026012</td>
                    <td><a href="...">81617447</a></td>
                    <td>3-7-2026</td>
                    <td>JAKARTA</td><td>JAKARTA</td><td>MAKASSAR</td><td>MAKASSAR</td>
                    <td>SATRIA REG</td><td>KREDIT</td><td>1</td><td>1</td><td>CGK539398263</td>
                    <td>POD - DELIVERED</td>
                    <td>Max SLA : 4 days Shipping Duration : 3 days Status : POD - DELIVERED IN SLA</td>
                    <td></td><td>CGK07251279A</td>;<td>UPGH106260724</td>
                    <td><a href="...">Riwayat</a></td>
                </tr>
            </tbody>
        </table>
    """

    def test_extract_sla_info(self):
        detail = "Max SLA : 4 days Shipping Duration : 3 days Status : POD - DELIVERED IN SLA"
        max_sla, dur, status = extract_sla_info(detail)
        self.assertEqual(max_sla, "4 days")
        self.assertEqual(dur, "3 days")
        self.assertEqual(status, "POD - DELIVERED IN SLA")

    def test_parse_tracking_focus_html(self):
        records = parse_tracking_focus_html(self.SAMPLE_HTML)
        self.assertEqual(len(records), 1)
        rec = records[0]
        self.assertEqual(rec["no"], 1)
        self.assertEqual(rec["awb_no"], "CGK8161744700066")
        self.assertEqual(rec["reference_no"], "NONPO03072026012")
        self.assertEqual(rec["destination"], "MAKASSAR")
        self.assertEqual(rec["max_sla"], "4 days")
        self.assertEqual(rec["shipping_duration"], "3 days")
        self.assertEqual(rec["sla_status"], "POD - DELIVERED IN SLA")
        self.assertEqual(rec["pod_by"], "UPGH106260724")

    def test_read_and_export_tracking_focus(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            wb_in = Workbook()
            ws_in = wb_in.active
            ws_in.append(["No. Referance"])
            ws_in.append(["NONPO03072026012"])
            ws_in.append(["NONPO03072026999"])
            in_path = Path(temp_dir) / "test_input.xlsx"
            wb_in.save(in_path)

            with open(in_path, "rb") as fp:
                targets = read_reference_targets_workbook(fp)
            self.assertEqual(targets, ["NONPO03072026012", "NONPO03072026999"])

            records = parse_tracking_focus_html(self.SAMPLE_HTML)
            out_path = Path(temp_dir) / "output.xlsx"
            export_tracking_focus(records, targets, out_path, search_by="a.reference_no")

            wb_out = load_workbook(out_path)
            ws_out = wb_out.active
            self.assertEqual(ws_out.title, "Riwayat Tracking Focus")
            self.assertEqual(ws_out.max_row, 3)  # Header + 1 found + 1 not found
            self.assertEqual(ws_out.cell(row=2, column=2).value, "NONPO03072026012")
            self.assertEqual(ws_out.cell(row=2, column=3).value, "CGK8161744700066")
            self.assertEqual(ws_out.cell(row=3, column=2).value, "NONPO03072026999")
            self.assertEqual(ws_out.cell(row=3, column=15).value, "TIDAK DITEMUKAN")


class TrackingHistoryPastedInputTests(unittest.TestCase):
    def test_milestones_without_tlc(self):
        records = [
            {"process": "ENTRI VERIFIED", "datetime": "2026-06-23 13:31:27", "location": "KANTOR PUSAT / CGK1"},
            {"process": "OUTGOING SMU", "datetime": "2026-06-25 06:48:35", "location": "KANTOR PUSAT / CGK1"},
            {"process": "INCOMING SMU", "datetime": "2026-06-26 03:04:33", "location": "SURABAYA / SUB1"},
            {"process": "POD", "datetime": "2026-07-03 09:31:00", "location": "PALANGKARAYA / PKYD103184915"},
        ]
        verified, outgoing, incoming, pod = tracking_milestones(records, "")
        self.assertIsNotNone(verified)
        self.assertIsNotNone(outgoing)
        self.assertEqual(incoming.strftime("%Y-%m-%d %H:%M:%S"), "2026-06-26 03:04:33")
        self.assertEqual(pod.strftime("%Y-%m-%d %H:%M:%S"), "2026-07-03 09:31:00")

    def test_api_create_job_with_pasted_awb_text(self):
        import json
        with app.test_client() as client:
            c = CoresysClient()
            c.username = "TEST_USER"
            with client.session_transaction() as sess:
                sess["session_id"] = "test-sess-paste"
            with clients_lock:
                clients["test-sess-paste"] = c

            payload = {
                "workflow": "tracking_history",
                "awb_text": "CGK10001\nCGK10002 PKY",
                "tracking_mode": "milestone",
                "delay_seconds": 0,
                "parallelism": 1,
            }
            res = client.post("/api/jobs", data=json.dumps(payload), content_type="application/json")
            self.assertEqual(res.status_code, 201)
            job = res.json["job"]
            self.assertEqual(job["workflow"], "tracking_history")
            batches = job["batches"]
            self.assertEqual(len(batches), 1)
            self.assertEqual(batches[0]["item_total"], 2)
            self.assertIn("2 AWB", batches[0]["label"])


if __name__ == "__main__":
    unittest.main()
