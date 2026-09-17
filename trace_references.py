import re
import sys
import time
from html.parser import HTMLParser
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from coresys import CoresysClient, BASE_URL


class TableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows = []
        self.headers = []
        self._current_row = None
        self._current_cell = None
        self._is_header = False

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._current_row = []
        elif tag == "th":
            self._is_header = True
            self._current_cell = []
        elif tag == "td":
            self._is_header = False
            self._current_cell = []
        elif tag == "br" and self._current_cell is not None:
            self._current_cell.append("\n")

    def handle_data(self, data):
        if self._current_cell is not None:
            self._current_cell.append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._current_cell is not None:
            text = " ".join("".join(self._current_cell).split())
            if self._is_header:
                self.headers.append(text)
            elif self._current_row is not None:
                self._current_row.append(text)
            self._current_cell = None
        elif tag == "tr":
            if self._current_row and any(self._current_row):
                self.rows.append(self._current_row)
            self._current_row = None


def extract_sla_info(detail_status_text: str):
    m1 = re.search(r"Max SLA\s*:\s*(.*?)(?=\s*Shipping Duration|\s*Status|$)", detail_status_text, re.I)
    m2 = re.search(r"Shipping Duration\s*:\s*(.*?)(?=\s*Status|$)", detail_status_text, re.I)
    m3 = re.search(r"Status\s*:\s*(.*)$", detail_status_text, re.I)

    max_sla = m1.group(1).strip() if m1 else ""
    duration = m2.group(1).strip() if m2 else ""
    sla_status = m3.group(1).strip() if m3 else ""
    return max_sla, duration, sla_status


def main():
    excel_path = Path(r"d:\SAPX\Data Analisis\SAP DOWNLOADER\TRACING.xlsx")
    if not excel_path.exists():
        print(f"File {excel_path} tidak ditemukan!")
        sys.exit(1)

    wb_in = openpyxl.load_workbook(excel_path)
    ws_in = wb_in.active

    ref_list = []
    for row in ws_in.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            val = str(row[0]).strip()
            if val and val != "None":
                ref_list.append(val)

    print(f"Ditemukan {len(ref_list)} No. Referance dari {excel_path.name}")

    print("Menghubungkan ke CORESYS...")
    client = CoresysClient()
    username = os.environ.get("CORESYS_USER") or input("Username CORESYS: ").strip()
    import getpass
    password = os.environ.get("CORESYS_PASS") or getpass.getpass("Password: ").strip()
    pin = os.environ.get("CORESYS_PIN") or getpass.getpass("PIN: ").strip()
    login_info = client.login(username, password, pin)
    print(f"Login sukses sebagai: {login_info['username']}")

    batch_size = 200
    all_parsed_rows = []

    for i in range(0, len(ref_list), batch_size):
        batch = ref_list[i : i + batch_size]
        batch_no = (i // batch_size) + 1
        total_batches = (len(ref_list) + batch_size - 1) // batch_size
        print(f"Mengambil data batch {batch_no}/{total_batches} ({len(batch)} referensi)...")

        resp = client.session.post(
            f"{BASE_URL}/tracking/show_list_data_tracking/",
            data={
                "val[]": batch,
                "key": "a.reference_no",
                "key_rowstate": "0",
            },
            headers={
                "X-Requested-With": "XMLHttpRequest",
                "Referer": f"{BASE_URL}/tracking/focus",
            },
            timeout=90,
        )
        resp.raise_for_status()

        parser = TableParser()
        parser.feed(resp.text)
        print(f"  -> Batch {batch_no} selesai: {len(parser.rows)} baris data diterima.")
        all_parsed_rows.extend(parser.rows)
        time.sleep(0.5)

    print(f"\nTotal baris data riwayat diperoleh: {len(all_parsed_rows)}")

    output_headers = [
        "No",
        "No. Referance",
        "No. AWB",
        "No. Master",
        "Tanggal",
        "Asal",
        "District Asal",
        "Tujuan",
        "Kota Tujuan",
        "Jenis Layanan",
        "Transaksi",
        "Kilo",
        "Koli",
        "No Invoice",
        "Status",
        "Detail Status",
        "Max SLA",
        "Shipping Duration",
        "SLA Status",
        "No. Resi Pickup",
        "Dibuat Oleh",
        "Di POD Oleh",
    ]

    ref_to_rows = {}
    for r in all_parsed_rows:
        if len(r) >= 19:
            ref_no = r[2].strip()
            ref_to_rows.setdefault(ref_no, []).append(r)

    final_data = []
    counter = 1
    found_refs = set()

    for ref in ref_list:
        if ref in ref_to_rows:
            found_refs.add(ref)
            for r in ref_to_rows[ref]:
                awb_no = r[1].replace(";", "").strip()
                ref_no = r[2].strip()
                master_no = r[3].strip()
                tanggal = r[4].strip()
                asal = r[5].strip()
                dist_asal = r[6].strip()
                tujuan = r[7].strip()
                kota_tujuan = r[8].strip()
                layanan = r[9].strip()
                transaksi = r[10].strip()
                kilo = r[11].strip()
                koli = r[12].strip()
                invoice = r[13].strip()
                status = r[14].strip()
                detail_status = r[15].strip()
                resi_pickup = r[16].replace(";", "").strip()
                dibuat_oleh = r[17].replace(";", "").strip()
                pod_oleh = r[18].replace(";", "").strip()

                max_sla, duration, sla_status = extract_sla_info(detail_status)

                final_data.append([
                    counter,
                    ref_no,
                    awb_no,
                    master_no,
                    tanggal,
                    asal,
                    dist_asal,
                    tujuan,
                    kota_tujuan,
                    layanan,
                    transaksi,
                    kilo,
                    koli,
                    invoice,
                    status,
                    detail_status,
                    max_sla,
                    duration,
                    sla_status,
                    resi_pickup,
                    dibuat_oleh,
                    pod_oleh,
                ])
                counter += 1
        else:
            final_data.append([
                counter,
                ref,
                "-",
                "-",
                "-",
                "-",
                "-",
                "-",
                "-",
                "-",
                "-",
                "-",
                "-",
                "-",
                "TIDAK DITEMUKAN",
                "-",
                "-",
                "-",
                "-",
                "-",
                "-",
                "-",
            ])
            counter += 1

    print(f"Referensi ditemukan: {len(found_refs)} / {len(ref_list)}")
    if len(found_refs) < len(ref_list):
        print(f"Referensi tidak ditemukan: {len(ref_list) - len(found_refs)}")

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Hasil Riwayat Tracing"

    header_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    ws_out.append(output_headers)
    for col_idx in range(1, len(output_headers) + 1):
        cell = ws_out.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_font = Font(name="Segoe UI", size=9)
    for row in final_data:
        ws_out.append(row)
        curr_row = ws_out.max_row
        for col_idx in range(1, len(row) + 1):
            cell = ws_out.cell(row=curr_row, column=col_idx)
            cell.font = row_font
            cell.border = thin_border
            if col_idx in (1, 5, 11, 12, 13, 17, 18):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")

    for col in ws_out.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_out.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    output_file = Path(r"d:\SAPX\Data Analisis\SAP DOWNLOADER\TRACING_RESULT.xlsx")
    wb_out.save(output_file)
    print(f"\nBerhasil disimpan ke: {output_file}")


if __name__ == "__main__":
    main()
