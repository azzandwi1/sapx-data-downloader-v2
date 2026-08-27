from __future__ import annotations

from concurrent.futures import Future, ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
import re
import threading
import time
from typing import Any, BinaryIO, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from coresys import CoresysClient, CoresysError


HEADER_FILL = PatternFill("solid", fgColor="174A5B")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FILL = PatternFill("solid", fgColor="174A5B")
TITLE_FONT = Font(color="FFFFFF", bold=True, size=16)


def _styled_row(sheet, values, *, header: bool = False, title: bool = False):
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        if header:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        elif title:
            cell.fill = TITLE_FILL
            cell.font = TITLE_FONT
        cells.append(cell)
    return cells


def _portal_datetime(value: str) -> datetime | str:
    try:
        return datetime.strptime(value, "%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError):
        return value


def _configure_sheet(sheet, widths: list[int], freeze: str) -> None:
    sheet.freeze_panes = freeze
    sheet.sheet_view.showGridLines = False
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _excel_text(value: object) -> object:
    if not isinstance(value, str):
        return value
    value = value[:32767]
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _event_at_tlc(record: dict[str, Any], tlc: str) -> bool:
    if not tlc:
        return False
    location = str(record.get("location", "")).upper()
    operator = location.rsplit("/", 1)[-1].strip()
    return operator.startswith(tlc.upper())


def tracking_milestones(records: list[dict[str, Any]], tlc: str) -> tuple[object, object, object, object]:
    verified = next(
        (_portal_datetime(row["datetime"]) for row in records if row["process"].upper() == "ENTRI VERIFIED"),
        None,
    )
    outgoing = next(
        (_portal_datetime(row["datetime"]) for row in records if row["process"].upper() == "OUTGOING SMU"),
        None,
    )
    incoming = next(
        (_portal_datetime(row["datetime"]) for row in records
         if row["process"].upper() == "INCOMING SMU" and _event_at_tlc(row, tlc)),
        None,
    )
    pod = next(
        (_portal_datetime(row["datetime"]) for row in records
         if row["process"].upper() == "POD" and _event_at_tlc(row, tlc)),
        None,
    )
    return verified, outgoing, incoming, pod


def first_pod_courier(records: list[dict[str, Any]]) -> str | None:
    pod = next((row for row in records if str(row.get("process", "")).upper() == "POD"), None)
    if not pod:
        return None
    location = str(pod.get("location", ""))
    if "/" not in location:
        return None
    return location.rsplit("/", 1)[1].strip() or None


def _note_field(note: object, field: str) -> str | None:
    match = re.search(rf"\[{re.escape(field)}:\s*([^\]]*)\]", str(note), re.I)
    return match.group(1).strip() or None if match else None


def pickup_attempt_summary(records: list[dict[str, Any]]) -> list[object]:
    request = next(
        (row for row in records if str(row.get("process", "")).upper() == "ENTRI (SEDANG DI PICKUP)"),
        None,
    )
    attempts = [
        row for row in records if str(row.get("process", "")).upper() == "ENTRI (PENDING PICKUP)"
    ]
    first_attempt = attempts[0] if attempts else None
    picked_up = next(
        (row for row in records if str(row.get("process", "")).upper() == "PICKED UP"),
        None,
    )
    not_ready_attempts = [
        row for row in attempts
        if re.search(r"\b(?:PAKET\s+)?BELUM\s+(?:READY|SIAP)\b|\bNOT\s+READY\b", str(row.get("note", "")), re.I)
    ]

    request_at = _portal_datetime(request.get("datetime", "")) if request else None
    attempt_at = _portal_datetime(first_attempt.get("datetime", "")) if first_attempt else None
    sla = None
    late_pickup = False
    if isinstance(request_at, datetime):
        if isinstance(attempt_at, datetime):
            days = (attempt_at.date() - request_at.date()).days
            late_pickup = days > 1
            sla = f"LATE PICKUP (H+{days})" if late_pickup else f"SESUAI (H+{days})"
        else:
            deadline = request_at.date().toordinal() + 1
            late_pickup = datetime.now().date().toordinal() > deadline
            sla = "LATE PICKUP - BELUM ADA PERCOBAAN" if late_pickup else "MENUNGGU BATAS H+1"

    if not_ready_attempts:
        status = "SUDAH ADA PERCOBAAN PICKUP - PAKET BELUM READY"
    elif attempts:
        status = "SUDAH ADA PERCOBAAN PICKUP - ALASAN LAIN"
    elif picked_up:
        status = "PICKUP BERHASIL - TANPA RIWAYAT PENDING"
    elif request:
        status = "BELUM ADA HASIL PERCOBAAN"
    else:
        status = "DATA PICKUP TIDAK DITEMUKAN"
    if late_pickup:
        status = f"LATE PICKUP - {status}"

    return [
        _portal_datetime(request.get("datetime", "")) if request else None,
        _portal_datetime(first_attempt.get("datetime", "")) if first_attempt else None,
        _note_field(first_attempt.get("note", ""), "KURIR") if first_attempt else None,
        _note_field(first_attempt.get("note", ""), "KETERANGAN") if first_attempt else None,
        len(attempts),
        len(not_ready_attempts),
        _portal_datetime(picked_up.get("datetime", "")) if picked_up else None,
        sla,
        status,
    ]


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_awb_targets_workbook(source: BinaryIO, require_tlc: bool = True) -> list[dict[str, str]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        header_row = None
        awb_column = None
        tlc_column = None
        for row_number, row in enumerate(sheet.iter_rows(max_row=25, values_only=True), start=1):
            labels = [re.sub(r"[^A-Z0-9]+", " ", _cell_text(value).upper()).strip() for value in row]
            candidate_awb = next((index for index, label in enumerate(labels) if "AWB" in label), None)
            candidate_tlc = next((index for index, label in enumerate(labels) if "TLC" in label), None)
            if candidate_awb is not None and (candidate_tlc is not None or not require_tlc):
                header_row = row_number
                awb_column = candidate_awb
                tlc_column = candidate_tlc
                break
        if header_row is None or awb_column is None or (require_tlc and tlc_column is None):
            expected = "Kolom 'No. AWB' dan 'TLC Tujuan'" if require_tlc else "Kolom 'No. AWB'"
            raise ValueError(f"{expected} tidak ditemukan pada 25 baris pertama.")

        seen: set[tuple[str, str]] = set()
        targets: list[dict[str, str]] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            awb = _cell_text(row[awb_column] if awb_column < len(row) else None).upper()
            tlc = _cell_text(
                row[tlc_column] if tlc_column is not None and tlc_column < len(row) else None
            ).upper()
            if not awb:
                continue
            if not re.fullmatch(r"[A-Z0-9._/-]+", awb):
                raise ValueError(f"No. AWB tidak valid: {awb[:80]}")
            if require_tlc and not re.fullmatch(r"[A-Z0-9]{2,10}", tlc):
                raise ValueError(f"TLC Tujuan tidak valid untuk AWB {awb}: {tlc or '(kosong)'}")
            key = (awb, tlc)
            if key not in seen:
                seen.add(key)
                targets.append({"awb": awb, "tlc": tlc})
        if not targets:
            raise ValueError("File Excel tidak memiliki baris AWB yang dapat diproses.")
        return targets
    finally:
        workbook.close()


def export_tracking_history(
    client: CoresysClient,
    awb_targets: list[dict[str, str] | str],
    destination: Path,
    parallelism: int = 2,
    delay_seconds: int = 0,
    include_summary: bool = True,
    include_history: bool = True,
    tracking_mode: str = "milestone",
    progress: Callable[[int, int], None] | None = None,
    cancelled: Callable[[], bool] | None = None,
) -> tuple[Path, int, int, int]:
    if not awb_targets:
        raise ValueError("Masukkan minimal satu nomor AWB.")
    if tracking_mode not in {"milestone", "courier_pod", "pickup_attempt"}:
        raise ValueError("Jenis data trace & tracking tidak dikenal.")
    targets = [
        {"awb": str(item.get("awb", "")).strip(), "tlc": str(item.get("tlc", "")).strip().upper()}
        if isinstance(item, dict) else {"awb": str(item).strip(), "tlc": ""}
        for item in awb_targets
    ]

    workbook = Workbook(write_only=True)
    if tracking_mode == "courier_pod":
        result_sheet = workbook.create_sheet("POD Pertama")
        result_widths = [23, 28]
        result_headers = ["No. AWB", "ID Kurir POD Pertama"]
    elif tracking_mode == "pickup_attempt":
        result_sheet = workbook.create_sheet("Verifikasi Pickup")
        result_widths = [23, 24, 27, 26, 48, 20, 23, 24, 20, 38]
        result_headers = [
            "No. AWB", "Tanggal Permintaan Pickup", "Percobaan Pickup Pertama",
            "Kurir Percobaan Pertama", "Alasan Percobaan Pertama", "Jumlah Percobaan",
            "Jumlah Belum Ready", "Tanggal Pickup Berhasil", "SLA Percobaan Pertama", "Status Verifikasi",
        ]
    else:
        result_sheet = workbook.create_sheet("Milestone Tujuan")
        result_widths = [23, 22, 31, 38, 34]
        result_headers = [
            "No. AWB", "Tgl. Verifikasi", "Tanggal Outgoing SMU Pertama",
            "Tanggal Incoming SMU Pertama (TLC Tujuan)", "Tanggal POD Pertama (TLC Tujuan)",
        ]
    _configure_sheet(result_sheet, result_widths, "A2")
    result_sheet.append(_styled_row(result_sheet, result_headers, header=True))

    def empty_result(awb: str, status: str) -> list[object]:
        row = [_excel_text(awb), *([None] * (len(result_headers) - 1))]
        if tracking_mode == "pickup_attempt":
            row[-1] = _excel_text(status)
        return row

    summary = workbook.create_sheet("Ringkasan") if include_summary else None
    if summary:
        _configure_sheet(summary, [23, 17, 22, 22, 31, 43, 34], "A5")
    history_sheets: list[dict[str, object]] = []

    def add_history_sheet():
        number = len(history_sheets) + 1
        name = "Semua History" if number == 1 else f"Semua History {number}"
        sheet = workbook.create_sheet(name)
        _configure_sheet(sheet, [23, 8, 27, 27, 27, 22, 43, 75], "A2")
        sheet.append(_styled_row(sheet, [
            "AWB", "No.", "Proses", "No. Dokumen", "Referensi / Tag",
            "Tanggal & Jam", "Lokasi / Oleh", "Keterangan",
        ], header=True))
        state = {"sheet": sheet, "rows": 1}
        history_sheets.append(state)
        return state

    history_state = add_history_sheet() if include_history else None

    if summary:
        summary.append(_styled_row(summary, ["RINGKASAN TRACE & TRACKING", "", "", "", "", "", ""], title=True))
        summary.append(["Sumber", "https://online.coresyssap.com/tracking/focus", "", "", "", "", ""])
        summary.append([])
        summary.append(_styled_row(summary, [
            "AWB", "Total Aktivitas", "Aktivitas Pertama", "Aktivitas Terakhir",
            "Proses Terakhir", "Lokasi / Oleh Terakhir", "Status Ekstraksi",
        ], header=True))
    found = 0
    missing = 0
    failed = 0
    completed = 0
    milestones: dict[int, list[object]] = {}

    pace_lock = threading.Lock()
    next_request_at = [0.0]

    def fetch(awb: str):
        if delay_seconds:
            with pace_lock:
                now = time.monotonic()
                wait_seconds = max(0.0, next_request_at[0] - now)
                next_request_at[0] = max(now, next_request_at[0]) + delay_seconds
            if wait_seconds:
                time.sleep(wait_seconds)
        return client.fork().fetch_tracking_history(awb)

    executor = ThreadPoolExecutor(max_workers=max(1, min(parallelism, 12)), thread_name_prefix="tracking-history")
    futures: dict[Future, tuple[int, dict[str, str]]] = {
        executor.submit(fetch, target["awb"]): (index, target)
        for index, target in enumerate(targets)
    }
    try:
        for future in as_completed(futures):
            if cancelled and cancelled():
                for pending in futures:
                    pending.cancel()
                raise CoresysError("Pekerjaan dibatalkan.")

            index, target = futures[future]
            awb = target["awb"]
            tlc = target["tlc"]
            try:
                records = future.result()
                if records:
                    found += 1
                    first = records[0]
                    last = records[-1]
                    if tracking_mode == "courier_pod":
                        milestones[index] = [_excel_text(awb), _excel_text(first_pod_courier(records))]
                    elif tracking_mode == "pickup_attempt":
                        milestones[index] = [_excel_text(awb), *map(_excel_text, pickup_attempt_summary(records))]
                    else:
                        verified, outgoing, incoming, pod = tracking_milestones(records, tlc)
                        milestones[index] = [_excel_text(awb), verified, outgoing, incoming, pod]
                    if summary:
                        summary.append([
                            _excel_text(awb), len(records), _portal_datetime(first["datetime"]),
                            _portal_datetime(last["datetime"]), _excel_text(last["process"]),
                            _excel_text(last["location"]), "BERHASIL",
                        ])
                    if include_history:
                        for record in records:
                            if history_state["rows"] >= 1_048_576:
                                history_state = add_history_sheet()
                            history_state["sheet"].append([
                                _excel_text(awb), record["no"], _excel_text(record["process"]),
                                _excel_text(record["document"]), _excel_text(record["reference"]),
                                _portal_datetime(record["datetime"]), _excel_text(record["location"]),
                                _excel_text(record["note"]),
                            ])
                            history_state["rows"] += 1
                else:
                    missing += 1
                    milestones[index] = empty_result(awb, "AWB TIDAK DITEMUKAN")
                    if summary:
                        summary.append([_excel_text(awb), 0, None, None, None, None, "TIDAK DITEMUKAN"])
            except Exception as exc:
                failed += 1
                milestones[index] = empty_result(awb, f"GAGAL: {str(exc)[:240]}")
                if summary:
                    summary.append([_excel_text(awb), 0, None, None, None, None, _excel_text(f"GAGAL: {str(exc)[:240]}")])
            completed += 1
            if progress:
                progress(completed, len(targets))
    finally:
        executor.shutdown(wait=True, cancel_futures=True)

    for index in range(len(targets)):
        result_sheet.append(milestones[index])
    last_column = get_column_letter(len(result_headers))
    result_sheet.auto_filter.ref = f"A1:{last_column}{len(targets) + 1}"
    if summary:
        summary.auto_filter.ref = f"A4:G{len(targets) + 4}"
    for state in history_sheets:
        state["sheet"].auto_filter.ref = f"A1:H{state['rows']}"
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination, found, missing, failed
